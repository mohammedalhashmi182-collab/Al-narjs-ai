"""CSV/XLSX uploads for the data analyst: size-limited, parsed into a compact summary
(columns, row count, numeric stats, preview). Raw files are not stored."""

from __future__ import annotations

import csv
import io
import os
import statistics

ALLOWED = ("csv", "xlsx")
MAX_UPLOAD_BYTES = int(os.getenv("AGENT_UPLOAD_MAX_BYTES", str(5 * 1024 * 1024)))
MAX_ROWS = 50_000
MAX_COLS = 100
PREVIEW_ROWS = 10


class UploadError(ValueError):
    pass


def kind_for(filename: str) -> str:
    ext = (filename or "").rsplit(".", 1)[-1].lower()
    if ext not in ALLOWED:
        raise UploadError("نوع الملف غير مدعوم — CSV أو XLSX فقط")
    return ext


def _rows_csv(data: bytes) -> list[list[str]]:
    for enc in ("utf-8-sig", "cp1256", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    out = []
    for i, row in enumerate(csv.reader(io.StringIO(text), dialect)):
        if i > MAX_ROWS:
            break
        out.append(row[:MAX_COLS])
    return out


def _rows_xlsx(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise UploadError("تعذّر قراءة ملف Excel") from exc
    ws = wb.worksheets[0]
    out = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i > MAX_ROWS:
            break
        out.append(["" if v is None else str(v) for v in row[:MAX_COLS]])
    wb.close()
    return out


def _num(v: str):
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def parse(filename: str, data: bytes) -> dict:
    kind = kind_for(filename)
    if len(data) > MAX_UPLOAD_BYTES:
        raise UploadError(f"الملف أكبر من الحد المسموح ({MAX_UPLOAD_BYTES // (1024 * 1024)} ميجابايت)")
    if not data:
        raise UploadError("الملف فارغ")
    rows = _rows_csv(data) if kind == "csv" else _rows_xlsx(data)
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        raise UploadError("لا توجد بيانات في الملف")
    header = [str(h).strip() or f"col{i + 1}" for i, h in enumerate(rows[0])]
    body = rows[1:]
    columns = []
    for i, name in enumerate(header):
        vals = [r[i] for r in body if i < len(r) and str(r[i]).strip() != ""]
        nums = [n for n in (_num(v) for v in vals) if n is not None]
        col = {"name": name, "non_empty": len(vals)}
        if vals and len(nums) >= 0.8 * len(vals):
            col.update(type="number", min=min(nums), max=max(nums), sum=round(sum(nums), 2),
                       mean=round(statistics.fmean(nums), 2))
        else:
            col.update(type="text", distinct=len(set(vals)), top=sorted(set(vals), key=vals.count, reverse=True)[:5])
        columns.append(col)
    return {"kind": kind, "rows": len(body), "columns": columns,
            "preview": [dict(zip(header, r)) for r in body[:PREVIEW_ROWS]]}


def context_block(files: list[dict]) -> str:
    if not files:
        return ""
    parts = ["\n\n[ملفات بيانات رفعها العميل — حلّلها واستند إلى أرقامها]:"]
    for f in files:
        s = f.get("summary") or {}
        parts.append(f"• {f.get('filename')} — {s.get('rows', 0)} صف")
        for c in s.get("columns", [])[:30]:
            if c.get("type") == "number":
                parts.append(f"  - {c['name']}: رقمي، أدنى {c['min']}، أعلى {c['max']}، متوسط {c['mean']}، مجموع {c['sum']}")
            else:
                parts.append(f"  - {c['name']}: نصي، {c.get('distinct', 0)} قيمة مختلفة، الأكثر: {', '.join(map(str, c.get('top', [])))}")
        for row in s.get("preview", [])[:5]:
            parts.append("  > " + "; ".join(f"{k}={v}" for k, v in list(row.items())[:12]))
    return "\n".join(parts)
